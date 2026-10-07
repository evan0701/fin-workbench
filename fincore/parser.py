"""mapping.yaml → 标准科目余额模型。

mapping.yaml 是整个体系的"协议"：用户只需描述自己家 ERP 导出文件的样子
（哪一行是表头、哪几列是金额、科目编码和名称怎么拆），引擎就能把它转换成
标准模型。换软件、换格式，只换 mapping，不换引擎。
"""

from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

import yaml
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from .model import AnchorResult, StandardTB, TBRow

_TOL = 0.01  # 金额比较容差（元）


def load_mapping(path: str | Path) -> dict:
    """读取并做最小校验。字段不全时给出可读的报错，方便用户自查 mapping。"""
    raw = yaml.safe_load(Path(path).read_text(encoding="utf-8"))
    for key in ("source", "columns"):
        if key not in raw:
            raise ValueError(f"mapping 缺少必需段: {key}")
    cols = raw["columns"]
    if "code_name" not in cols and "code" not in cols:
        raise ValueError("columns 里要么给 code_name（编码名称同列），要么给 code + name 分列")
    return raw


def _col_idx(col_letter: str) -> int:
    return column_index_from_string(col_letter.upper())


def _num(value) -> float:
    """容忍金蝶导出里的空串、全角空格、千分位字符串。"""
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("\u3000", "").replace(",", "").replace("¥", "").replace("￥", "")
    if text in ("", "-", "—"):
        return 0.0
    return float(text)


def _parse_code_name(text: str, split_regex: str | None) -> tuple[str, str]:
    """处理"1002　银行存款"这类编码名称挤在一列的情况。"""
    text = str(text).strip()
    if not split_regex:
        return text, ""
    m = re.match(split_regex, text)
    if not m:
        raise ValueError(f"科目列无法按 mapping 的 split_regex 拆分: {text!r}")
    groups = m.groups()
    return groups[0].strip(), (groups[1].strip() if len(groups) > 1 else "")


def _level_of(code: str, widths: list[int] | None = None) -> int:
    """科目级次推断。

    默认按 2 位分段（4位=1级，6位=2级，8位=3级）。部分账套用 3 位分段
    （4/7/10 位），在 mapping 里配 level_widths: [4, 7, 10] 覆盖；
    编码长度不在 widths 里时退回默认规则。
    """
    if widths:
        try:
            return widths.index(len(code)) + 1
        except ValueError:
            pass
    return max(1, len(code) // 2 - 1)


def parse(mapping: dict, source_file: str | Path, prev_model: StandardTB | None = None) -> tuple[StandardTB, list[AnchorResult]]:
    """按 mapping 解析导出文件 → (标准模型, 锚点校验结果)。"""
    src_cfg = dict(mapping["source"])
    src_cfg["file"] = str(source_file)  # CLI --file 覆盖 mapping 里的默认文件
    sheet_name = src_cfg.get("sheet")
    header_row = int(src_cfg.get("header_row", 1))
    skip_patterns = [re.compile(p) for p in src_cfg.get("skip_patterns", ["合计", "制表"])]
    unit_scale = float(mapping.get("unit_scale", 1.0))
    level_widths = mapping.get("level_widths")

    cols = mapping["columns"]
    code_name_cfg = cols.get("code_name")
    code_cfg, name_cfg = cols.get("code"), cols.get("name")
    split_regex = (code_name_cfg or {}).get("split_regex") or (code_cfg or {}).get("split_regex")

    opening_cfg, period_cfg, closing_cfg = cols["opening"], cols["period"], cols["closing"]

    wb = load_workbook(source_file, read_only=True, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb[wb.sheetnames[0]]
    sheet_used = ws.title

    rows: list[TBRow] = []
    for row_no, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if row_no <= header_row:
            continue
        cells = list(row)
        first = cells[0] if cells else None
        first_text = str(first).strip() if first is not None else ""
        if not first_text:
            continue
        if any(p.search(first_text) for p in skip_patterns):
            continue

        if code_name_cfg is not None:
            raw_text = str(first).replace("\u3000", " ").strip()
            code, name = _parse_code_name(raw_text, split_regex)
        else:
            code = str(cells[_col_idx(code_cfg["col"]) - 1]).strip()
            name = str(cells[_col_idx(name_cfg["col"]) - 1]).strip() if name_cfg else ""
        if not code or not re.match(r"^\d+(\.\d+)*$", code):
            continue  # 表尾说明行等非科目行

        def get(cfg_key: str, sub: str) -> float:
            letter = cols[cfg_key][sub]
            return _num(cells[_col_idx(letter) - 1]) * unit_scale

        rows.append(
            TBRow(
                code=code,
                name=name,
                level=_level_of(code, level_widths),
                opening_debit=get("opening", "debit"),
                opening_credit=get("opening", "credit"),
                period_debit=get("period", "debit"),
                period_credit=get("period", "credit"),
                closing_debit=get("closing", "debit"),
                closing_credit=get("closing", "credit"),
                src={"file": Path(source_file).name, "sheet": sheet_used, "row": row_no},
            )
        )
    wb.close()

    if not rows:
        raise ValueError(f"没解析到任何科目行，请检查 mapping 的 header_row({header_row}) 和列配置: {source_file}")

    if mapping.get("leaf_only"):
        # 金蝶等导出会把一级/二级/末级科目的全额逐级重复展示，全量读取会双重计算。
        # leaf_only 只保留"编码不是任何其他编码前缀"的末级科目行。
        codes = [r.code for r in rows]
        rows = [r for r in rows if not any(o != r.code and o.startswith(r.code) for o in codes)]
        if not rows:
            raise ValueError("leaf_only 过滤后没有剩余科目行，请检查编码列是否正确")

    meta = {
        "source_file": Path(source_file).name,
        "sheet": sheet_used,
        "period": src_cfg.get("period", ""),
        "company": src_cfg.get("company", ""),
        "unit": mapping.get("unit_note", "元"),
        "mapping_name": mapping.get("name", ""),
        "mapping_version": mapping.get("version", 1),
        "leaf_only": bool(mapping.get("leaf_only")),
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M"),
    }
    model = StandardTB(meta=meta, rows=rows)
    anchors = run_anchors(model, prev_model, requested=mapping.get("anchors", []))
    return model, anchors


def run_anchors(model: StandardTB, prev_model: StandardTB | None = None, requested: list[str] | None = None) -> list[AnchorResult]:
    """锚点校验。全绿才允许进入下游（视图/填写），这是"可审计"的第一道闸。"""
    requested = requested or ["opening_balance", "closing_balance", "codes_unique", "carry_forward"]
    results: list[AnchorResult] = []
    t = model.totals()

    if "opening_balance" in requested:
        diff = t["opening_debit"] - t["opening_credit"]
        results.append(
            AnchorResult("opening_balance", "PASS" if abs(diff) <= _TOL else "FAIL",
                         f"期初借方合计 {t['opening_debit']:.2f} vs 期初贷方合计 {t['opening_credit']:.2f}，差额 {diff:+.2f}")
        )
    if "closing_balance" in requested:
        diff = t["closing_debit"] - t["closing_credit"]
        results.append(
            AnchorResult("closing_balance", "PASS" if abs(diff) <= _TOL else "FAIL",
                         f"期末借方合计 {t['closing_debit']:.2f} vs 期末贷方合计 {t['closing_credit']:.2f}，差额 {diff:+.2f}")
        )
    if "codes_unique" in requested:
        codes = [r.code for r in model.rows]
        dup = sorted({c for c in codes if codes.count(c) > 1})
        empty = [r.code for r in model.rows if not r.code.strip()]
        ok = not dup and not empty
        detail = "科目编码无重复、无空值" if ok else f"重复编码: {dup} / 空编码 {len(empty)} 行"
        results.append(AnchorResult("codes_unique", "PASS" if ok else "FAIL", detail))
    if "carry_forward" in requested:
        if prev_model is None:
            results.append(AnchorResult("carry_forward", "SKIPPED", "未提供上月模型，跳过衔接校验"))
        else:
            prev_close = {r.code: r for r in prev_model.rows}
            bad = []
            for r in model.rows:
                p = prev_close.get(r.code)
                if p is None:
                    if r.is_balance_subject:
                        bad.append(f"{r.code} 上月不存在")
                    continue
                if abs(r.net_opening - p.net_closing) > _TOL:
                    bad.append(f"{r.code} 期初{r.net_opening:.2f} ≠ 上月期末{p.net_closing:.2f}")
            results.append(
                AnchorResult("carry_forward", "PASS" if not bad else "FAIL",
                             "与上月期末衔接一致" if not bad else "；".join(bad[:5]) + (f"（共{len(bad)}处）" if len(bad) > 5 else ""))
            )
    return results
