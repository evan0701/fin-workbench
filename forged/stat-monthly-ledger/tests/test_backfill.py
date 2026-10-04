"""回填考试：新铸技能必须用上期人工版当考卷，全对才上岗。

本测试就是 fill-forge 第五步的自动化形态：
用 2025-08 数据源回填，对照 demo/统计台账模板/统计月度台账_2025-08_人工版.xlsx，
任何一格差异 > 0.005 万元即失败。
"""

from __future__ import annotations

import sys
from pathlib import Path

import pytest

SKILL_DIR = Path(__file__).resolve().parents[1]          # forged/stat-monthly-ledger
REPO = Path(__file__).resolve().parents[3]               # fin-workbench 根
sys.path.insert(0, str(SKILL_DIR))

from fill import main as fill_main  # noqa: E402


@pytest.fixture(scope="module", autouse=True)
def _ensure_demo_sources():
    """demo 产物缺失时自动重新生成（生成器只用星澜虚构数据）。"""
    sources = REPO / "demo" / "示例导出"
    if not (sources / "科目余额表_2025-08_乱格式.xlsx").exists():
        import subprocess

        subprocess.run(
            [sys.executable, str(REPO / "demo" / "gen_demo_source.py")],
            check=True, cwd=REPO,
        )


def test_backfill_exam_matches_manual_version(tmp_path):
    out = tmp_path / "回填考试_2025-08.xlsx"
    rc = fill_main(
        [
            "--source", str(REPO / "demo" / "示例导出" / "科目余额表_2025-08_乱格式.xlsx"),
            "--period", "2025-08",
            "--out", str(out),
            "--check", str(REPO / "demo" / "统计台账模板" / "统计月度台账_2025-08_人工版.xlsx"),
        ]
    )
    assert rc == 0, "回填考试失败：技能填写结果与人工版存在差异"
    assert out.exists()
    assert (tmp_path / "回填考试_2025-08.xlsx.复核清单.md").exists()


def test_fill_refuses_when_anchor_fails(tmp_path):
    """铁律 1：锚点不全绿不填报。构造一个借贷不平的源文件，fill 必须拒绝。"""
    import openpyxl

    bad = tmp_path / "bad_source.xlsx"
    src = REPO / "demo" / "示例导出" / "科目余额表_2025-08_乱格式.xlsx"
    wb = openpyxl.load_workbook(src)
    ws = wb.active
    # 找到 1002 银行存款所在行，把期末借方加 100 元破坏平衡
    for row in range(5, ws.max_row + 1):
        if ws.cell(row=row, column=1).value and str(ws.cell(row=row, column=1).value).startswith("1002"):
            ws.cell(row=row, column=6, value=float(ws.cell(row=row, column=6).value) + 100.0)
            break
    wb.save(bad)

    rc = fill_main(
        [
            "--source", str(bad),
            "--period", "2025-08",
            "--out", str(tmp_path / "should_not_exist.xlsx"),
        ]
    )
    assert rc == 2, "锚点 FAIL 时必须拒绝填报（退出码 2）"
    assert not (tmp_path / "should_not_exist.xlsx").exists()
